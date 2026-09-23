"""Release verifier for the browser-generated OOXML workbook fixture."""

from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook


def main() -> int:
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "artifacts/release/xlsx-smoke.xlsx")
    required_parts = {
        "[Content_Types].xml",
        "_rels/.rels",
        "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
        "xl/styles.xml",
    }

    with zipfile.ZipFile(target) as archive:
        bad_file = archive.testzip()
        if bad_file is not None:
            raise RuntimeError(f"ZIP CRC failure: {bad_file}")
        names = set(archive.namelist())
        missing = sorted(required_parts - names)
        if missing:
            raise RuntimeError(f"Missing OOXML parts: {missing}")
        xml_parts = [name for name in names if name.endswith((".xml", ".rels"))]
        for name in xml_parts:
            ElementTree.fromstring(archive.read(name))

    workbook = load_workbook(target, read_only=True, data_only=False)
    personnel_mode = target.name == "personnel-list.xlsx"
    expected_sheets = {"人員合格名單"} if personnel_mode else {"QR-28-01", "QR-28-03", "觀察事項", "建議追蹤", "稽核前準備"}
    missing_sheets = sorted(expected_sheets - set(workbook.sheetnames))
    if missing_sheets:
        raise RuntimeError(f"Missing worksheets: {missing_sheets}")
    first_value = workbook["人員合格名單"]["A1"].value if personnel_mode else workbook["QR-28-01"]["A1"].value
    expected_title = "姓名" if personnel_mode else "年度內部稽核計畫"
    if not isinstance(first_value, str) or expected_title not in first_value:
        raise RuntimeError("Expected Chinese workbook title was not preserved")

    result = {
        "file": str(target.resolve()),
        "zip_crc": "PASS",
        "xml_parts": len(xml_parts),
        "openpyxl_load": "PASS",
        "sheet_count": len(workbook.sheetnames),
        "sheet_names": workbook.sheetnames,
        "cjk_title": "PASS",
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    workbook.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
